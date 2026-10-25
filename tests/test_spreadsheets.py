"""XLSX -> CSV through the real registry and runner (workbooks are generated)."""

import csv
import datetime
import io
import threading
import zipfile

import openpyxl
import pytest

from fileconverter.conversion import runner
from fileconverter.core.errors import InvalidInputError, InvalidOptionError, UnsupportedFormatError
from fileconverter.core.models import ConversionRequest, JobStatus
from helpers import temp_files

BOM = "﻿"


def make_workbook(path, sheets: dict[str, list[list]]):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    book.save(path)
    return path


def set_cached_value(path, sheet_file: str, value: str):
    """Give formula cells a saved result, as Excel would (openpyxl never computes one)."""
    with zipfile.ZipFile(path) as archive:
        items = {item.filename: archive.read(item) for item in archive.infolist()}
    items[sheet_file] = items[sheet_file].replace(b"<v />", f"<v>{value}</v>".encode())
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in items.items():
            archive.writestr(name, data)


def convert(registry, source, **options):
    request = ConversionRequest(source, "csv", source.parent, options)
    return runner.run_job(request, registry)


def read(path):
    return path.read_bytes().decode("utf-8")


def outputs(folder):
    return sorted(p.name for p in folder.iterdir() if p.suffix == ".csv")


def test_single_sheet_uses_plain_name_and_documented_dialect(registry, tmp_path):
    source = make_workbook(
        tmp_path / "book.xlsx",
        {"Only": [["Name", "Note"], ["Ёлка, Größe", 'say "hi"'], ["multi\nline", 1.5]]},
    )
    original = source.read_bytes()

    result = convert(registry, source, delimiter="comma", decimal="dot")

    assert result.status is JobStatus.SUCCEEDED, result.error
    assert [p.name for p in result.output_paths] == ["book.csv"]
    assert read(result.output_path) == (
        BOM + 'Name,Note\r\n"Ёлка, Größe","say ""hi"""\r\n"multi\nline",1.5\r\n'
    )
    assert source.read_bytes() == original


def test_every_worksheet_becomes_one_csv_in_order(registry, tmp_path):
    source = make_workbook(
        tmp_path / "book.xlsx",
        {"Первый": [["a"]], "Second": [["b"]], "Третий лист": [["c"]]},
    )

    result = convert(registry, source)

    assert [p.name for p in result.output_paths] == [
        "book_Первый.csv",
        "book_Second.csv",
        "book_Третий лист.csv",
    ]
    assert [read(p) for p in result.output_paths] == [BOM + "a\r\n", BOM + "b\r\n", BOM + "c\r\n"]
    assert temp_files(tmp_path) == []


def test_empty_worksheet_gives_empty_csv_and_warning(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"Data": [["x"]], "Blank": []})

    result = convert(registry, source)

    assert result.status is JobStatus.SUCCEEDED
    blank = tmp_path / "book_Blank.csv"
    assert blank.read_bytes() == b""
    assert any("'Blank' is empty" in w for w in result.warnings)


def test_sheet_names_are_sanitized_and_deduplicated(registry, tmp_path):
    source = make_workbook(
        tmp_path / "book.xlsx", {"a<b": [[1]], "a>b": [[2]], "CON": [[3]], 'q"t|x': [[4]]}
    )

    result = convert(registry, source)

    assert [p.name for p in result.output_paths] == [
        "book_a_b.csv",
        "book_a_b_2.csv",
        "book_CON_.csv",  # reserved device name on Windows
        "book_q_t_x.csv",
    ]


@pytest.mark.parametrize("choice", ["2", "Zwei", "zwei"])
def test_one_worksheet_can_be_selected(registry, tmp_path, choice):
    source = make_workbook(tmp_path / "book.xlsx", {"Eins": [[1]], "Zwei": [[2]], "Drei": [[3]]})

    result = convert(registry, source, sheet=choice)

    assert [p.name for p in result.output_paths] == ["book_Zwei.csv"]
    assert read(result.output_path) == BOM + "2\r\n"


def test_unknown_worksheet_lists_available_ones(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"Eins": [[1]], "Zwei": [[2]]})

    result = convert(registry, source, sheet="7")

    assert isinstance(result.error, InvalidOptionError)
    assert "1: Eins, 2: Zwei" in str(result.error)
    assert outputs(tmp_path) == []


def test_formula_without_saved_result_exports_formula_and_warns(registry, tmp_path):
    source = make_workbook(tmp_path / "f.xlsx", {"S": [[1, 2, "=A1+B1"]]})

    result = convert(registry, source, delimiter="comma", formulas="raw")

    assert read(result.output_path) == BOM + "1,2,=A1+B1\r\n"
    assert any("1 formula cell(s) had no saved result" in w for w in result.warnings)


def test_formula_with_saved_result_exports_the_value(registry, tmp_path):
    source = make_workbook(tmp_path / "f.xlsx", {"S": [[1, 2, "=A1+B1"]]})
    set_cached_value(source, "xl/worksheets/sheet1.xml", "3")

    result = convert(registry, source, delimiter="comma")

    assert read(result.output_path) == BOM + "1,2,3\r\n"
    assert not any("formula" in w for w in result.warnings)


def test_value_formatting(registry, tmp_path):
    row = [
        True,
        False,
        datetime.datetime(2024, 1, 31),
        datetime.datetime(2024, 1, 31, 13, 45, 7),
        0.1,
        42,
        None,
        "end",
    ]
    source = make_workbook(tmp_path / "v.xlsx", {"S": [row]})

    result = convert(registry, source, delimiter="comma", decimal="dot")

    assert read(result.output_path) == (
        BOM + "TRUE,FALSE,2024-01-31,2024-01-31 13:45:07,0.1,42,,end\r\n"
    )


def test_rows_stay_rectangular_with_gaps(registry, tmp_path):
    book = openpyxl.Workbook()
    book.active["A1"] = "a"
    book.active["C1"] = "c"
    book.active["B3"] = "b"
    source = tmp_path / "gaps.xlsx"
    book.save(source)

    result = convert(registry, source, delimiter="comma")

    assert read(result.output_path) == BOM + "a,,c\r\n,,\r\n,b,\r\n"


TRICKY_ROWS = [["Ёлka; Größe", "a,b", "tab\there"], ['say "hi"', "multi\nline", -1.5]]


def test_default_dialect_is_semicolon_with_bom(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"S": TRICKY_ROWS})

    result = convert(registry, source)  # no options: the defaults GUI and CLI also use

    assert result.output_path.read_bytes().startswith(b"\xef\xbb\xbf")
    assert read(result.output_path) == (
        BOM + '"Ёлka; Größe";a,b;tab\there\r\n"say ""hi""";"multi\nline";-1,5\r\n'
    )


@pytest.mark.parametrize(
    ("delimiter", "expected"),
    [
        ("semicolon", '"Ёлka; Größe";a,b;tab\there\r\n"say ""hi""";"multi\nline";-1,5\r\n'),
        ("comma", 'Ёлka; Größe,"a,b",tab\there\r\n"say ""hi""","multi\nline","-1,5"\r\n'),
        ("tab", 'Ёлka; Größe\ta,b\t"tab\there"\r\n"say ""hi"""\t"multi\nline"\t-1,5\r\n'),
    ],
)
def test_each_delimiter_quotes_what_it_must(registry, tmp_path, delimiter, expected):
    source = make_workbook(tmp_path / "book.xlsx", {"S": TRICKY_ROWS})

    result = convert(registry, source, delimiter=delimiter)

    assert read(result.output_path) == BOM + expected


def test_delimiter_applies_to_every_worksheet(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"A": [[1, 2]], "B": [[3, 4]], "C": []})

    result = convert(registry, source, delimiter="tab")

    assert [read(p) for p in result.output_paths] == [BOM + "1\t2\r\n", BOM + "3\t4\r\n", ""]


def test_unknown_delimiter_is_rejected(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"S": [[1]]})
    assert isinstance(convert(registry, source, delimiter="pipe").error, InvalidOptionError)
    assert outputs(tmp_path) == []


# --- decimal separator ----------------------------------------------------------

NUMBERS = [1.5, -1.5, 42, -7, 0, 0.1, 1234567.891, 0.000123, 1.5e20, -2.5e-07]


def test_numbers_use_decimal_comma_by_default(registry, tmp_path):
    source = make_workbook(tmp_path / "n.xlsx", {"S": [NUMBERS]})

    result = convert(registry, source)  # defaults: semicolon + decimal comma

    assert read(result.output_path) == (
        BOM + "1,5;-1,5;42;-7;0;0,1;1234567,891;0,000123;1,5e+20;-2,5e-07\r\n"
    )


def test_decimal_dot_keeps_python_number_text(registry, tmp_path):
    source = make_workbook(tmp_path / "n.xlsx", {"S": [NUMBERS]})

    result = convert(registry, source, decimal="dot")

    assert read(result.output_path) == (
        BOM + "1.5;-1.5;42;-7;0;0.1;1234567.891;0.000123;1.5e+20;-2.5e-07\r\n"
    )


def test_full_precision_survives_the_comma(registry, tmp_path):
    value = 0.1 + 0.2  # 0.30000000000000004: must not be rounded away
    source = make_workbook(tmp_path / "n.xlsx", {"S": [[0.5]]})
    # openpyxl itself rounds floats when writing, so store the exact value in the XML,
    # as Excel would.
    with zipfile.ZipFile(source) as archive:
        items = {item.filename: archive.read(item) for item in archive.infolist()}
    sheet = "xl/worksheets/sheet1.xml"
    items[sheet] = items[sheet].replace(b"<v>0.5</v>", f"<v>{value!r}</v>".encode())
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in items.items():
            archive.writestr(name, data)

    text = read(convert(registry, source).output_path)[1:].strip()

    assert text == "0,30000000000000004"
    assert float(text.replace(",", ".")) == value


@pytest.mark.parametrize("decimal", ["comma", "dot"])
def test_text_dates_and_formulas_are_never_changed(registry, tmp_path, decimal):
    row = [
        "Product 1.5",
        "001.50",
        "1,5",
        "1.5",  # number-looking text stays text: it is stored as a string
        "v1.2.3",
        datetime.datetime(2024, 1, 31, 13, 45, 7, 500000),
        datetime.time(8, 30, 15, 250000),
        "=A1*1.5",  # formula without saved result: its text is exported
    ]
    source = make_workbook(tmp_path / "t.xlsx", {"S": [row]})

    result = convert(registry, source, decimal=decimal, formulas="raw")

    assert read(result.output_path) == (
        BOM + "Product 1.5;001.50;1,5;1.5;v1.2.3;2024-01-31 13:45:07.500000;"
        "08:30:15.250000;=A1*1.5\r\n"
    )


def test_saved_formula_results_are_numbers(registry, tmp_path):
    source = make_workbook(tmp_path / "f.xlsx", {"S": [[3, "=A1/4"]]})
    set_cached_value(source, "xl/worksheets/sheet1.xml", "0.75")

    assert read(convert(registry, source).output_path) == BOM + "3;0,75\r\n"


def test_decimal_comma_with_comma_delimiter_is_quoted(registry, tmp_path):
    source = make_workbook(
        tmp_path / "q.xlsx", {"S": [[1.5, "a;b", 'say "hi"', "two\nlines", "x,y"]]}
    )

    result = convert(registry, source, delimiter="comma")

    text = read(result.output_path)
    assert text == BOM + '"1,5",a;b,"say ""hi""","two\nlines","x,y"\r\n'
    assert next(csv.reader(io.StringIO(text[1:]))) == [
        "1,5",
        "a;b",
        'say "hi"',
        "two\nlines",
        "x,y",
    ]


def test_semicolon_and_decimal_comma_for_german_excel(registry, tmp_path):
    source = make_workbook(
        tmp_path / "de.xlsx", {"S": [["Preis", "Menge"], [12.99, 3], ["Größe; Maß", -0.5]]}
    )

    text = read(convert(registry, source, delimiter="semicolon", decimal="comma").output_path)

    assert text == BOM + 'Preis;Menge\r\n12,99;3\r\n"Größe; Maß";-0,5\r\n'
    rows = list(csv.reader(io.StringIO(text[1:]), delimiter=";"))
    assert rows == [["Preis", "Menge"], ["12,99", "3"], ["Größe; Maß", "-0,5"]]


def test_every_worksheet_uses_the_same_decimal_separator(registry, tmp_path):
    source = make_workbook(tmp_path / "m.xlsx", {"A": [[1.25]], "B": [[-2.5]], "C": [[3]]})

    result = convert(registry, source, decimal="dot", delimiter="tab")

    assert [read(p) for p in result.output_paths] == [
        BOM + "1.25\r\n",
        BOM + "-2.5\r\n",
        BOM + "3\r\n",
    ]
    result = convert(registry, source)
    assert [read(p) for p in result.output_paths] == [
        BOM + "1,25\r\n",
        BOM + "-2,5\r\n",
        BOM + "3\r\n",
    ]


@pytest.mark.parametrize("mode", ["protect", "raw"])
def test_decimal_comma_does_not_interfere_with_formula_protection(registry, tmp_path, mode):
    source = make_workbook(tmp_path / "p.xlsx", {"S": [[-1.5, "-1,5", "=1+1", "-x"]]})

    text = read(convert(registry, source, formulas=mode).output_path)

    if mode == "protect":  # numbers and number-like text untouched, formula text protected
        assert text == BOM + "-1,5;-1,5;'=1+1;'-x\r\n"
    else:
        assert text == BOM + "-1,5;-1,5;=1+1;-x\r\n"


def test_unknown_decimal_separator_is_rejected(registry, tmp_path):
    source = make_workbook(tmp_path / "b.xlsx", {"S": [[1.5]]})
    assert isinstance(convert(registry, source, decimal="space").error, InvalidOptionError)
    assert outputs(tmp_path) == []


# --- formula injection --------------------------------------------------------

DANGEROUS = [
    '=HYPERLINK("http://x","y")',
    "+cmd|' /C calc'!A0",
    "-2+3",
    "@SUM(A1)",
    "  =1+1",  # leading spaces
    "\t=1+1",  # leading tab
    "\r=1+1",
    "＝1+1",  # full-width equals sign
    "- bullet point",
    "+49 30 123456",  # phone number: a spreadsheet would try to calculate it
]
HARMLESS = ["-5", "+1,5", "-0.25", "1e5", "plain text", "", "a=b", "e-mail@example.com"]


@pytest.mark.parametrize("text", DANGEROUS)
def test_protect_mode_prefixes_formula_like_text(registry, tmp_path, text):
    source = make_workbook(tmp_path / "book.xlsx", {"S": [[text]]})

    result = convert(registry, source, delimiter="comma")  # protect is the default

    cell = next(csv.reader(io.StringIO(read(result.output_path)[1:])))[0]
    # XML parsing turns a stored "\r" into "\n" (XML end-of-line normalization).
    assert cell == "'" + text.replace("\r", "\n")
    assert any("prefixed with '" in w for w in result.warnings)


@pytest.mark.parametrize("text", HARMLESS)
def test_protect_mode_leaves_harmless_text_alone(registry, tmp_path, text):
    source = make_workbook(tmp_path / "book.xlsx", {"S": [[text, "x"]]})

    result = convert(registry, source, delimiter="comma")

    assert next(csv.reader(io.StringIO(read(result.output_path)[1:])))[0] == text
    assert not any("prefixed" in w for w in result.warnings)


def test_numbers_are_never_prefixed(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"S": [[-5, -1.5, 0, 3]]})
    result = convert(registry, source, delimiter="comma")
    assert read(result.output_path) == BOM + '-5,"-1,5",0,3\r\n'  # quoted: it holds the delimiter


def test_unsaved_formula_is_exported_as_text_in_protect_mode(registry, tmp_path):
    source = make_workbook(tmp_path / "f.xlsx", {"S": [[1, 2, "=A1+B1"]]})
    result = convert(registry, source, delimiter="comma")
    assert read(result.output_path) == BOM + "1,2,'=A1+B1\r\n"


def test_raw_mode_writes_values_unchanged(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"S": [["=1+1", "@x", "-y"]]})
    result = convert(registry, source, delimiter="comma", formulas="raw")
    assert read(result.output_path) == BOM + "=1+1,@x,-y\r\n"
    assert not any("prefixed" in w for w in result.warnings)


def test_existing_outputs_are_never_overwritten(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"S1": [["new"]], "S2": [["new"]]})
    taken = tmp_path / "book_S1.csv"
    taken.write_text("user data", encoding="utf-8")

    result = convert(registry, source)

    assert [p.name for p in result.output_paths] == ["book_S1 (1).csv", "book_S2.csv"]
    assert taken.read_text(encoding="utf-8") == "user data"


def test_damaged_workbook(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {"S": [["x"]]})
    with zipfile.ZipFile(source) as archive:
        items = {item.filename: archive.read(item) for item in archive.infolist()}
    items["xl/workbook.xml"] = b"<not-xml"
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in items.items():
            archive.writestr(name, data)

    result = convert(registry, source)

    assert isinstance(result.error, InvalidInputError)
    assert outputs(tmp_path) == [] and temp_files(tmp_path) == []


def test_zip_without_workbook_is_not_xlsx(registry, tmp_path):
    fake = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(fake, "w") as archive:
        archive.writestr("data.txt", "x")
    assert isinstance(convert(registry, fake).error, UnsupportedFormatError)


def test_cancel_between_worksheets_leaves_nothing(registry, tmp_path):
    source = make_workbook(tmp_path / "book.xlsx", {f"S{i}": [[i]] for i in range(5)})
    cancel = threading.Event()

    def on_progress(fraction):
        if fraction and fraction < 1:
            cancel.set()  # after the first worksheet was written

    result = runner.run_job(
        ConversionRequest(source, "csv", tmp_path), registry, cancel, on_progress
    )

    assert result.status is JobStatus.CANCELLED
    assert outputs(tmp_path) == []
    assert temp_files(tmp_path) == []
